from io import BytesIO

from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook

from app.main import app

client = TestClient(app)


def test_descriptive_stats_csv_upload_create_task_and_validate_success() -> None:
    upload_response = client.post(
        "/api/uploads",
        files={"file": ("descriptive.csv", b"age,group\n23,A\n27,B\n31,A\n", "text/csv")},
    )

    assert upload_response.status_code == 200
    upload_payload = upload_response.json()
    assert upload_payload["success"] is True
    assert upload_payload["message"] == "uploaded"

    upload_id = upload_payload["data"]["upload_id"]

    task_response = client.post(
        "/api/tasks",
        json={"method_slug": "descriptive_stats", "upload_id": upload_id},
    )

    assert task_response.status_code == 200
    task_payload = task_response.json()
    assert task_payload["success"] is True
    assert task_payload["message"] == "task created"

    task_id = task_payload["data"]["task_id"]

    validate_response = client.post(f"/api/tasks/{task_id}/validate")

    assert validate_response.status_code == 200
    validate_payload = validate_response.json()
    assert validate_payload["success"] is True
    assert validate_payload["message"] == "validated"
    assert validate_payload["data"]["status"] == "validated"
    assert validate_payload["data"]["validation_passed"] is True
    assert validate_payload["data"]["validation"]["passed"] is True
    assert validate_payload["data"]["validation"]["summary"] == {"error_count": 0, "warning_count": 0}
    assert validate_payload["data"]["validation"]["stats"]["row_count"] == 3
    assert validate_payload["data"]["validation"]["stats"]["column_count"] == 2


def test_upload_supports_xlsx_files() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["age", "group"])
    worksheet.append([23, "A"])
    worksheet.append([27, "B"])

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    response = client.post(
        "/api/uploads",
        files={
            "file": (
                "descriptive.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["success"] is True
    assert payload["data"]["file_type"] == "xlsx"


def test_validate_returns_structured_field_errors() -> None:
    upload_response = client.post(
        "/api/uploads",
        files={"file": ("broken.csv", b",group\n23,A\n27,B\n", "text/csv")},
    )
    upload_id = upload_response.json()["data"]["upload_id"]

    task_response = client.post(
        "/api/tasks",
        json={"method_slug": "descriptive_stats", "upload_id": upload_id},
    )
    task_id = task_response.json()["data"]["task_id"]

    validate_response = client.post(f"/api/tasks/{task_id}/validate")

    assert validate_response.status_code == 200
    payload = validate_response.json()

    assert payload["success"] is True
    assert payload["data"]["validation_passed"] is False
    assert payload["data"]["validation"]["passed"] is False
    assert payload["data"]["validation"]["summary"]["error_count"] >= 1

    first_error = payload["data"]["validation"]["errors"][0]
    assert first_error["field"] == "column_1"
    assert "缺少列名" in first_error["message"]
    assert "补充明确的列名" in first_error["suggestion"]


def test_descriptive_stats_analysis_and_result_are_computed_from_uploaded_data() -> None:
    upload_response = client.post(
        "/api/uploads",
        files={"file": ("analysis.csv", b"age,score,group\n20,80,A\n30,90,B\n40,100,A\n", "text/csv")},
    )
    upload_id = upload_response.json()["data"]["upload_id"]

    task_response = client.post(
        "/api/tasks",
        json={"method_slug": "descriptive_stats", "upload_id": upload_id},
    )
    task_id = task_response.json()["data"]["task_id"]

    validate_response = client.post(f"/api/tasks/{task_id}/validate")
    assert validate_response.status_code == 200
    assert validate_response.json()["data"]["validation_passed"] is True

    analyze_response = client.post(f"/api/tasks/{task_id}/analyze")
    assert analyze_response.status_code == 200
    analyze_payload = analyze_response.json()
    assert analyze_payload["success"] is True
    assert analyze_payload["message"] == "analysis completed"
    assert analyze_payload["data"] == {"task_id": task_id, "status": "success"}

    result_response = client.get(f"/api/tasks/{task_id}/result")
    assert result_response.status_code == 200
    result_payload = result_response.json()

    assert result_payload["success"] is True
    assert result_payload["message"] == "ok"
    assert result_payload["data"]["task_id"] == task_id
    assert result_payload["data"]["method_slug"] == "descriptive_stats"
    assert result_payload["data"]["status"] == "success"

    numeric_table = next(table for table in result_payload["data"]["tables"] if table["key"] == "numeric_summary")
    assert numeric_table["columns"] == ["variable", "count", "mean", "median", "std_dev", "min", "max", "missing_count"]
    assert numeric_table["rows"][0] == ["age", 3, 30.0, 30.0, 10.0, 20.0, 40.0, 0]
    assert numeric_table["rows"][1] == ["score", 3, 90.0, 90.0, 10.0, 80.0, 100.0, 0]

    categorical_table = next(table for table in result_payload["data"]["tables"] if table["key"] == "categorical_frequency")
    assert categorical_table["rows"][0] == ["group", "A", 2, 0.6667]
    assert categorical_table["rows"][1] == ["group", "B", 1, 0.3333]

    chart = result_payload["data"]["charts"][0]
    assert chart["key"] == "numeric_means"
    assert chart["data"]["labels"] == ["age", "score"]
    assert chart["data"]["values"] == [30.0, 90.0]

    interpretation = result_payload["data"]["interpretation"]
    assert "30.00" in interpretation["plain_language"]
    assert "66.7%" in interpretation["academic_style"]


def test_regression_validation_returns_specific_field_errors() -> None:
    upload_response = client.post(
        "/api/uploads",
        files={"file": ("regression_invalid.csv", b"predictor_income,name\n10,A\n20,B\n", "text/csv")},
    )
    upload_id = upload_response.json()["data"]["upload_id"]

    task_response = client.post(
        "/api/tasks",
        json={"method_slug": "regression", "upload_id": upload_id},
    )
    task_id = task_response.json()["data"]["task_id"]

    validate_response = client.post(f"/api/tasks/{task_id}/validate")
    assert validate_response.status_code == 200

    payload = validate_response.json()
    assert payload["data"]["validation_passed"] is False

    fields = {issue["field"] for issue in payload["data"]["validation"]["errors"]}
    assert "outcome" in fields
    assert "predictors" in fields


def test_regression_analysis_and_docx_export_work_with_real_results() -> None:
    csv_content = (
        "outcome,predictor_income,predictor_age,control_experience\n"
        "29,10,1,5\n"
        "47,15,2,5\n"
        "67,20,3,5\n"
        "85,25,4,5\n"
        "105,30,5,5\n"
        "123,35,6,5\n"
    )
    upload_response = client.post(
        "/api/uploads",
        files={"file": ("regression.csv", csv_content.encode("utf-8"), "text/csv")},
    )
    upload_id = upload_response.json()["data"]["upload_id"]

    task_response = client.post(
        "/api/tasks",
        json={"method_slug": "regression", "upload_id": upload_id},
    )
    task_id = task_response.json()["data"]["task_id"]

    validate_response = client.post(f"/api/tasks/{task_id}/validate")
    assert validate_response.status_code == 200
    assert validate_response.json()["data"]["validation_passed"] is False

    corrected_csv = (
        "outcome,predictor_income,predictor_age,control_region\n"
        "24,10,1,1\n"
        "53,18,4,0\n"
        "51,22,2,1\n"
        "74,27,5,0\n"
        "72,31,3,1\n"
        "103,40,6,0\n"
    )
    corrected_upload_response = client.post(
        "/api/uploads",
        files={"file": ("regression_fixed.csv", corrected_csv.encode("utf-8"), "text/csv")},
    )
    corrected_upload_id = corrected_upload_response.json()["data"]["upload_id"]
    corrected_task_response = client.post(
        "/api/tasks",
        json={"method_slug": "regression", "upload_id": corrected_upload_id},
    )
    corrected_task_id = corrected_task_response.json()["data"]["task_id"]

    corrected_validate_response = client.post(f"/api/tasks/{corrected_task_id}/validate")
    assert corrected_validate_response.status_code == 200
    corrected_validate_payload = corrected_validate_response.json()
    assert corrected_validate_payload["data"]["validation_passed"] is True
    assert corrected_validate_payload["data"]["validation"]["summary"]["error_count"] == 0

    analyze_response = client.post(f"/api/tasks/{corrected_task_id}/analyze")
    assert analyze_response.status_code == 200
    assert analyze_response.json()["data"] == {"task_id": corrected_task_id, "status": "success"}

    result_response = client.get(f"/api/tasks/{corrected_task_id}/result")
    assert result_response.status_code == 200
    result_payload = result_response.json()["data"]

    assert result_payload["method_slug"] == "regression"
    assert result_payload["summary"]["title"] == "回归分析完成"

    model_summary = next(table for table in result_payload["tables"] if table["key"] == "model_summary")
    coefficient_table = next(table for table in result_payload["tables"] if table["key"] == "coefficients")
    fit_table = next(table for table in result_payload["tables"] if table["key"] == "fit_statistics")
    coefficient_chart = result_payload["charts"][0]

    assert model_summary["columns"] == ["metric", "value"]
    assert coefficient_table["columns"] == ["variable", "coef", "std_err", "t_value", "p_value", "significant"]
    assert fit_table["rows"][0][0] == "f_statistic"
    assert coefficient_chart["key"] == "coefficient_chart"
    assert "predictor_income" in coefficient_chart["data"]["labels"]
    assert "R²" in result_payload["interpretation"]["academic_style"]

    export_response = client.post(
        f"/api/tasks/{corrected_task_id}/export",
        json={"export_type": "docx"},
    )
    assert export_response.status_code == 200
    export_payload = export_response.json()
    assert export_payload["success"] is True
    assert export_payload["message"] == "export created"
    assert export_payload["data"]["download_url"].startswith("/api/exports/")

    download_response = client.get(export_payload["data"]["download_url"])
    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    document = Document(BytesIO(download_response.content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert any("回归分析报告" in paragraph for paragraph in paragraphs)
    assert any("结果摘要" in paragraph for paragraph in paragraphs)
